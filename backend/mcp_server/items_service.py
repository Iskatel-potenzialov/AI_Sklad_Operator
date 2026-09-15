"""
Items Service — работа с Excel-таблицей товаров склада канцелярии.
Аналог ExcelService, но для товаров вместо клиентов.
"""
import os
from pathlib import Path
from typing import List, Optional, Dict, Any
from datetime import datetime

from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter

from backend.mcp_server.schemas import (
    ItemData,
    GetItemsOutput,
    GetItemOutput,
    UpdateMinStockOutput,
)
from backend.mcp_server.validation import (
    validate_item_id,
    validate_excel_columns,
    REQUIRED_ITEM_COLUMNS,
)


# Путь к Excel-файлу товаров
ITEMS_FILE_PATH = os.getenv(
    "ITEMS_FILE_PATH",
    str(Path(__file__).parent.parent.parent / "data" / "items.xlsx")
)


class ItemsServiceError(Exception):
    """Ошибка Items-сервиса."""
    def __init__(self, message: str, error_code: str = "ITEMS_ERROR"):
        super().__init__(message)
        self.error_code = error_code


class ItemsService:
    """
    Сервис для работы с Excel-файлом товаров склада.
    Все операции с openpyxl инкапсулированы здесь.
    """
    
    def __init__(self, file_path: Optional[str] = None):
        self.file_path = Path(file_path or ITEMS_FILE_PATH)
        self._ensure_file_exists()
    
    def _ensure_file_exists(self):
        """Проверяет существование файла товаров."""
        if not self.file_path.exists():
            raise ItemsServiceError(
                f"Файл товаров не найден: {self.file_path}. "
                f"Запустите backend/scripts/generate_warehouse_data.py для создания тестовых данных.",
                error_code="FILE_NOT_FOUND"
            )
    
    def _load_workbook(self, data_only: bool = True):
        """Загружает workbook с обработкой ошибок."""
        if not self.file_path.exists():
            raise ItemsServiceError(
                f"Файл товаров не найден: {self.file_path}",
                error_code="FILE_NOT_FOUND"
            )
        try:
            return load_workbook(self.file_path, data_only=data_only)
        except Exception as e:
            raise ItemsServiceError(
                f"Не удалось открыть файл: {e}",
                error_code="FILE_READ_ERROR"
            )
    
    def _get_headers(self, ws) -> List[str]:
        """Получает заголовки из первой строки."""
        return [cell.value for cell in ws[1] if cell.value]
    
    def _validate_headers(self, ws):
        """Проверяет наличие обязательных колонок."""
        headers = self._get_headers(ws)
        ok, msg = validate_excel_columns(headers, REQUIRED_ITEM_COLUMNS)
        if not ok:
            raise ItemsServiceError(msg, error_code="INVALID_SCHEMA")
    
    def _find_item_row(self, ws, item_id: str) -> Optional[int]:
        """Находит номер строки товара по item_id."""
        headers = self._get_headers(ws)
        try:
            id_col_idx = headers.index("item_id") + 1
        except ValueError:
            raise ItemsServiceError(
                "Колонка 'item_id' не найдена в таблице",
                error_code="INVALID_SCHEMA"
            )
        
        for row_idx in range(2, ws.max_row + 1):
            if ws.cell(row=row_idx, column=id_col_idx).value == item_id:
                return row_idx
        return None
    
    def _row_to_item_data(self, ws, row_idx: int, headers: List[str]) -> ItemData:
        """Преобразует строку Excel в ItemData."""
        row_data = {}
        for col_idx, header in enumerate(headers, 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            row_data[header] = value if value is not None else ""
        
        # Приводим числовые поля
        try:
            quantity = int(row_data.get("quantity", 0) or 0)
        except (ValueError, TypeError):
            quantity = 0
        
        try:
            min_stock = int(row_data.get("min_stock", 0) or 0)
        except (ValueError, TypeError):
            min_stock = 0
        
        try:
            unit_price = float(row_data.get("unit_price", 0) or 0)
        except (ValueError, TypeError):
            unit_price = 0
        
        try:
            weight_kg = float(row_data.get("weight_kg", 0) or 0)
        except (ValueError, TypeError):
            weight_kg = 0
        
        try:
            volume_m3 = float(row_data.get("volume_m3", 0) or 0)
        except (ValueError, TypeError):
            volume_m3 = 0
        
        return ItemData(
            item_id=str(row_data.get("item_id", "")),
            product_name=str(row_data.get("product_name", "")),
            category=str(row_data.get("category", "")),
            quantity=quantity,
            min_stock=min_stock,
            supplier_id=str(row_data.get("supplier_id", "")),
            last_ordered=str(row_data.get("last_ordered", "") or ""),
            reorder_date=str(row_data.get("reorder_date", "") or ""),
            unit_price=unit_price,
            unit_of_measure=str(row_data.get("unit_of_measure", "")),
            weight_kg=weight_kg,
            volume_m3=volume_m3,
            location=str(row_data.get("location", "")),
            expiry_date=str(row_data.get("expiry_date", "") or ""),
            status=str(row_data.get("status", "")),
            ai_status=str(row_data.get("ai_status", "") or ""),
        )
    
    # ============================================================
    # Публичные методы (MCP-инструменты)
    # ============================================================
    
    def get_items(
        self,
        item_id: Optional[str] = None,
        category: Optional[str] = None,
    ) -> GetItemsOutput:
        """
        Получить список товаров.
        
        Если item_id указан — вернуть только этот товар.
        Если category указан — фильтровать по категории.
        """
        print(f"[ItemsService] get_items(item_id={item_id}, category={category})")
        
        try:
            wb = self._load_workbook()
            ws = wb.active
            self._validate_headers(ws)
            headers = self._get_headers(ws)
            
            items = []
            for row_idx in range(2, ws.max_row + 1):
                item = self._row_to_item_data(ws, row_idx, headers)
                
                # Пропускаем пустые строки
                if not item.item_id:
                    continue
                
                # Фильтр по item_id
                if item_id and item.item_id != item_id:
                    continue
                
                # Фильтр по категории
                if category and item.category != category:
                    continue
                
                items.append(item)
            
            wb.close()
            
            print(f"[ItemsService] get_items: found {len(items)} items")
            return GetItemsOutput(
                success=True,
                items=items,
                count=len(items),
            )
        
        except ItemsServiceError as e:
            print(f"[ItemsService] get_items error: {e}")
            return GetItemsOutput(
                success=False,
                error=str(e),
                error_code=e.error_code,
            )
        except Exception as e:
            print(f"[ItemsService] get_items unexpected error: {e}")
            return GetItemsOutput(
                success=False,
                error=f"Неожиданная ошибка: {str(e)}",
                error_code="UNEXPECTED_ERROR",
            )
    
    def get_item(self, item_id: str) -> GetItemOutput:
        """Получить данные конкретного товара."""
        print(f"[ItemsService] get_item(item_id={item_id})")
        
        # Валидация item_id
        ok, msg = validate_item_id(item_id)
        if not ok:
            return GetItemOutput(
                success=False,
                error=msg,
                error_code="INVALID_INPUT",
            )
        
        try:
            wb = self._load_workbook()
            ws = wb.active
            self._validate_headers(ws)
            headers = self._get_headers(ws)
            
            row_idx = self._find_item_row(ws, item_id)
            if not row_idx:
                wb.close()
                return GetItemOutput(
                    success=False,
                    error=f"Товар {item_id} не найден",
                    error_code="ITEM_NOT_FOUND",
                )
            
            item = self._row_to_item_data(ws, row_idx, headers)
            wb.close()
            
            print(f"[ItemsService] get_item: found {item.product_name}")
            return GetItemOutput(
                success=True,
                item=item,
            )
        
        except ItemsServiceError as e:
            return GetItemOutput(
                success=False,
                error=str(e),
                error_code=e.error_code,
            )
        except Exception as e:
            return GetItemOutput(
                success=False,
                error=f"Неожиданная ошибка: {str(e)}",
                error_code="UNEXPECTED_ERROR",
            )
    
    def update_min_stock(
        self,
        item_id: str,
        new_min_stock: int,
    ) -> UpdateMinStockOutput:
        """
        Изменить минимальный остаток товара.
        
        1. Проверяет существование товара
        2. Проверяет допустимость нового значения
        3. Читает актуальную запись перед изменением
        4. Выполняет изменение
        5. Сохраняет Excel
        6. Возвращает результат операции
        """
        print(f"[ItemsService] update_min_stock(item_id={item_id}, new_min_stock={new_min_stock})")
        
        # Валидация входных данных
        ok, msg = validate_item_id(item_id)
        if not ok:
            return UpdateMinStockOutput(
                success=False,
                error=msg,
                error_code="INVALID_INPUT",
            )
        
        if new_min_stock < 0:
            return UpdateMinStockOutput(
                success=False,
                error="Минимальный остаток не может быть отрицательным",
                error_code="INVALID_INPUT",
            )
        
        try:
            # Загружаем workbook для записи (data_only=False)
            wb = load_workbook(self.file_path)
            ws = wb.active
            self._validate_headers(ws)
            headers = self._get_headers(ws)
            
            # Ищем товар
            row_idx = self._find_item_row(ws, item_id)
            if not row_idx:
                wb.close()
                return UpdateMinStockOutput(
                    success=False,
                    error=f"Товар {item_id} не найден",
                    error_code="ITEM_NOT_FOUND",
                )
            
            # Читаем актуальное значение
            min_stock_col = headers.index("min_stock") + 1
            old_min_stock = ws.cell(row=row_idx, column=min_stock_col).value or 0
            
            # Выполняем изменение
            ws.cell(row=row_idx, column=min_stock_col).value = new_min_stock
            wb.save(self.file_path)
            wb.close()
            
            print(f"[ItemsService] update_min_stock: {old_min_stock} → {new_min_stock}")
            return UpdateMinStockOutput(
                success=True,
                item_id=item_id,
                changes={
                    "min_stock": {
                        "old": int(old_min_stock),
                        "new": new_min_stock,
                    }
                },
                message=f"Минимальный остаток товара {item_id} успешно изменён",
            )
        
        except ItemsServiceError as e:
            return UpdateMinStockOutput(
                success=False,
                error=str(e),
                error_code=e.error_code,
            )
        except Exception as e:
            return UpdateMinStockOutput(
                success=False,
                error=f"Неожиданная ошибка: {str(e)}",
                error_code="UNEXPECTED_ERROR",
            )


# Singleton
_items_service: Optional[ItemsService] = None


def get_items_service() -> ItemsService:
    """Получить экземпляр ItemsService."""
    global _items_service
    if _items_service is None:
        _items_service = ItemsService()
    return _items_service
