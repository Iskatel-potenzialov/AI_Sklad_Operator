"""
Suppliers Service — работа с Excel-таблицей поставщиков.
"""
import os
from pathlib import Path
from typing import List, Optional
from datetime import datetime

from openpyxl import load_workbook

from backend.mcp_server.schemas import (
    SupplierData,
    GetSuppliersOutput,
)


# Путь к Excel-файлу поставщиков
SUPPLIERS_FILE_PATH = os.getenv(
    "SUPPLIERS_FILE_PATH",
    str(Path(__file__).parent.parent.parent / "data" / "suppliers.xlsx")
)


class SuppliersServiceError(Exception):
    """Ошибка Suppliers-сервиса."""
    def __init__(self, message: str, error_code: str = "SUPPLIERS_ERROR"):
        super().__init__(message)
        self.error_code = error_code


class SuppliersService:
    """
    Сервис для работы с Excel-файлом поставщиков.
    """
    
    def __init__(self, file_path: Optional[str] = None):
        self.file_path = Path(file_path or SUPPLIERS_FILE_PATH)
        self._ensure_file_exists()
    
    def _ensure_file_exists(self):
        """Проверяет существование файла поставщиков."""
        if not self.file_path.exists():
            raise SuppliersServiceError(
                f"Файл поставщиков не найден: {self.file_path}. "
                f"Запустите backend/scripts/generate_warehouse_data.py для создания тестовых данных.",
                error_code="FILE_NOT_FOUND"
            )
    
    def _load_workbook(self, data_only: bool = True):
        """Загружает workbook с обработкой ошибок."""
        if not self.file_path.exists():
            raise SuppliersServiceError(
                f"Файл поставщиков не найден: {self.file_path}",
                error_code="FILE_NOT_FOUND"
            )
        try:
            return load_workbook(self.file_path, data_only=data_only)
        except Exception as e:
            raise SuppliersServiceError(
                f"Не удалось открыть файл: {e}",
                error_code="FILE_READ_ERROR"
            )
    
    def _get_headers(self, ws) -> List[str]:
        """Получает заголовки из первой строки."""
        return [cell.value for cell in ws[1] if cell.value]
    
    def _row_to_supplier_data(self, ws, row_idx: int, headers: List[str]) -> SupplierData:
        """Преобразует строку Excel в SupplierData."""
        row_data = {}
        for col_idx, header in enumerate(headers, 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            row_data[header] = value if value is not None else ""
        
        # Приводим числовые поля
        try:
            rating = float(row_data.get("rating", 0) or 0)
        except (ValueError, TypeError):
            rating = 0
        
        return SupplierData(
            supplier_id=str(row_data.get("supplier_id", "")),
            company_name=str(row_data.get("company_name", "")),
            contact_person=str(row_data.get("contact_person", "")),
            phone=str(row_data.get("phone", "")),
            email=str(row_data.get("email", "")),
            contract_start=str(row_data.get("contract_start", "") or ""),
            contract_end=str(row_data.get("contract_end", "") or ""),
            payment_terms=str(row_data.get("payment_terms", "")),
            rating=rating,
            last_contact=str(row_data.get("last_contact", "") or ""),
            status=str(row_data.get("status", "")),
        )
    
    def get_suppliers(self) -> GetSuppliersOutput:
        """
        Получить список всех поставщиков.
        """
        print(f"[SuppliersService] get_suppliers()")
        
        try:
            wb = self._load_workbook()
            ws = wb.active
            headers = self._get_headers(ws)
            
            suppliers = []
            for row_idx in range(2, ws.max_row + 1):
                supplier = self._row_to_supplier_data(ws, row_idx, headers)
                
                # Пропускаем пустые строки
                if not supplier.supplier_id:
                    continue
                
                suppliers.append(supplier)
            
            wb.close()
            
            print(f"[SuppliersService] get_suppliers: found {len(suppliers)} suppliers")
            return GetSuppliersOutput(
                success=True,
                suppliers=suppliers,
                count=len(suppliers),
            )
        
        except SuppliersServiceError as e:
            print(f"[SuppliersService] get_suppliers error: {e}")
            return GetSuppliersOutput(
                success=False,
                error=str(e),
                error_code=e.error_code,
            )
        except Exception as e:
            print(f"[SuppliersService] get_suppliers unexpected error: {e}")
            return GetSuppliersOutput(
                success=False,
                error=f"Неожиданная ошибка: {str(e)}",
                error_code="UNEXPECTED_ERROR",
            )


# Singleton
_suppliers_service = None


def get_suppliers_service() -> SuppliersService:
    """Получить экземпляр SuppliersService."""
    global _suppliers_service
    if _suppliers_service is None:
        _suppliers_service = SuppliersService()
    return _suppliers_service
